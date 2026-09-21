# -*- coding: utf-8 -*-
"""Регресс-тесты воркеров на реальных исходниках (fixtures/) и end-to-end
прогон основного скрипта. Эталонные числа — из выгрузки конвертера v112
от 31.08.2026 (папка 2026.09.17), с учётом согласованных отличий формата."""
import csv
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))
FIXTURES = BASE / "fixtures"

# Фикстуры (реальные прайсы) не публикуются в открытый репозиторий —
# на машинах без них регресс-тесты пропускаются, unit-тесты работают.
pytestmark = pytest.mark.skipif(
    not (FIXTURES / "Getsy.xlsx").exists(),
    reason="нет локальных фикстур (162 МБ реальных прайсов, вне публичного репо)",
)

from workers import getsy, treolan  # noqa: E402

E_NOTATION = re.compile(r"^[0-9,+-]+[eE][+-]?[0-9]+$")


def read_csv(path: Path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    return rows[0], rows[1:]


NUMERIC_COLS = (6, 7, 8, 9, 18)  # Ширина, Высота, Глубина, Вес, Объём


def assert_contract(header, rows):
    assert len(header) == 21
    assert header[0] == "№" and header[-1] == "Ошибка"
    for r in rows:
        assert len(r) == 21
        for i in NUMERIC_COLS:
            assert not E_NOTATION.match(r[i]), f"научная нотация: {r[i]}"


# ---------------- Getsy ----------------

@pytest.fixture(scope="module")
def getsy_out(tmp_path_factory):
    out = tmp_path_factory.mktemp("getsy") / "Getsy.csv"
    res = getsy.process(FIXTURES / "Getsy.xlsx", out, 1)
    return out, res


def test_getsy_rows_match_v112(getsy_out):
    out, res = getsy_out
    header, rows = read_csv(out)
    assert_contract(header, rows)
    assert len(rows) == 34337              # столько же, сколько у v112 на этом файле
    assert res.last_number == 34337

def test_getsy_first_row_values(getsy_out):
    out, _ = getsy_out
    _, rows = read_csv(out)
    r = rows[0]
    assert r[1] == "44000001847"           # Код = Номенклатурный номер
    assert r[2] == "RA598"                 # Артикул = Каталожный номер
    assert r[5].startswith("Комплектующие для ПК -> Корпуса")
    assert r[13] == "3+"                   # Количество: первое непустое из складов
    assert r[14] == "179,28"               # цена из фикстуры 31.08
    assert r[15] == "RUB"
    assert r[16] == "шт"
    assert r[17] == "Стандарт"

def test_getsy_dimensions_in_cm(getsy_out):
    """Метры исходника переводятся в сантиметры: точечная сверка по коду товара."""
    import openpyxl
    wb = openpyxl.load_workbook(FIXTURES / "Getsy.xlsx", read_only=True, data_only=True)
    ws = wb["Наличие и цены"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    i_code, i_depth = hdr.index("Номенклатурный номер"), hdr.index("Глубина (м)")
    src = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = row[i_depth]
        if isinstance(d, (int, float)) and d > 0:
            src = (str(row[i_code]), float(d))
            break
    wb.close()
    assert src, "в исходнике нет ненулевой глубины"
    out, _ = getsy_out
    _, rows = read_csv(out)
    got = next(r[8] for r in rows if r[1] == src[0])
    assert abs(float(got.replace(",", ".")) - src[1] * 100) < 1e-9  # м -> см

def test_getsy_quality_and_vat(getsy_out):
    out, _ = getsy_out
    _, rows = read_csv(out)
    assert {r[17] for r in rows} <= {"Стандарт", "Дефект упаковки", "Дефект товара",
                                     "Из ремонта", "Некомплект", "Демо", "Распродажа"}
    assert all(r[11] == "" for r in rows)  # НДС пуст (решение 17.09)


# ---------------- Treolan ----------------

@pytest.fixture(scope="module")
def treolan_out(tmp_path_factory):
    out = tmp_path_factory.mktemp("treolan") / "Treolan.csv"
    res = treolan.process(FIXTURES / "Treolan.xlsx", out, 1)
    return out, res


def test_treolan_rows_match_v112(treolan_out):
    out, res = treolan_out
    header, rows = read_csv(out)
    assert_contract(header, rows)
    assert len(rows) == 7868               # как у v112 на этом файле

def test_treolan_currency_split(treolan_out):
    out, _ = treolan_out
    _, rows = read_csv(out)
    cur = {"USD": 0, "RUB": 0}
    for r in rows:
        cur[r[15]] += 1
    assert cur == {"USD": 5147, "RUB": 2721}   # распределение как у v112

def test_treolan_first_item(treolan_out):
    out, _ = treolan_out
    _, rows = read_csv(out)
    r = rows[0]
    assert r[2] == "R260-01_noMEM"
    assert r[5].startswith("Серверы (Brand)")   # группа из строки-группы
    assert r[10] == "1 год"                     # гарантия прямым копированием
    assert r[13] == "<10"
    assert r[17] == "Стандарт"                  # без суффикса «(по имени файла)»
    assert r[8] == "78"                         # Глубина = «Длина [см]»
    assert r[7] == "25"                         # Высота

def test_treolan_nc_and_demo_quality(tmp_path):
    out_nc = tmp_path / "nc.csv"
    treolan.process(FIXTURES / "Treolan_NC.xlsx", out_nc, 1)
    _, rows = read_csv(out_nc)
    assert len(rows) == 291
    assert {r[17] for r in rows} == {"Некондиция"}

    out_demo = tmp_path / "demo.csv"
    treolan.process(FIXTURES / "Treolan_DEMO.xlsx", out_demo, 1)
    _, rows = read_csv(out_demo)
    assert len(rows) == 3
    assert {r[17] for r in rows} == {"Демо"}


def test_worker_fails_loudly_on_changed_format(tmp_path):
    """Смена структуры поставщиком должна давать внятную ошибку, а не сдвиг данных."""
    from common.reader import Sheet, SourceFormatError
    sh = Sheet(FIXTURES / "Treolan.xlsx", "Каталог", 3)
    with pytest.raises(SourceFormatError) as e:
        sh.col("Несуществующая колонка")
    assert "Несуществующая колонка" in str(e.value)


# ---------------- main.py: end-to-end ----------------

def test_main_end_to_end(tmp_path):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    archive_dir = tmp_path / "archive"
    input_dir.mkdir()
    for f in ("Getsy.xlsx", "Treolan.xlsx", "Treolan_NC.xlsx", "Treolan_DEMO.xlsx"):
        shutil.copy2(FIXTURES / f, input_dir / f)
    (input_dir / "Неизвестный прайс.xlsx").write_bytes(b"junk")

    settings = tmp_path / "settings.yaml"
    settings.write_text(f"""
dirs:
  input: {input_dir}
  output: {output_dir}
  archive: {archive_dir}
pricelists:
  - {{name: Getsy, match: "Getsy.xlsx", worker: getsy, result: "Getsy.csv"}}
  - {{name: Treolan, match: "Treolan.xlsx", worker: treolan, result: "Treolan.csv"}}
  - {{name: Treolan NC, match: "Treolan_NC.xlsx", worker: treolan, result: "Treolan_NC.csv"}}
  - {{name: Treolan DEMO, match: "Treolan_DEMO.xlsx", worker: treolan, result: "Treolan_DEMO.csv"}}
  - {{name: Выключенный, match: "Treolan_DEMO.xlsx", worker: treolan, result: "x.csv", enabled: false}}
""", encoding="utf-8")

    proc = subprocess.run([sys.executable, str(BASE / "main.py"), "--settings", str(settings)],
                          capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr

    # результаты и done.txt
    for f in ("Getsy.csv", "Treolan.csv", "Treolan_NC.csv", "Treolan_DEMO.csv", "done.txt"):
        assert (output_dir / f).exists(), f
    assert not (output_dir / "x.csv").exists()

    # сквозная нумерация «без зазоров» между файлами (ТЗ 4.1.5)
    _, g = read_csv(output_dir / "Getsy.csv")
    _, t = read_csv(output_dir / "Treolan.csv")
    _, nc = read_csv(output_dir / "Treolan_NC.csv")
    _, demo = read_csv(output_dir / "Treolan_DEMO.csv")
    assert int(t[0][0]) == int(g[-1][0]) + 1
    assert int(nc[0][0]) == int(t[-1][0]) + 1
    assert int(demo[0][0]) == int(nc[-1][0]) + 1

    # качество по имени файла должно переживать копирование во временный файл
    assert {r[17] for r in nc} == {"Некондиция"}
    assert {r[17] for r in demo} == {"Демо"}
    assert {r[17] for r in t} == {"Стандарт"}

    # оригиналы ушли в архивный подкаталог ГГГГ.ММ.ДД, вход пуст (кроме junk)
    subdirs = [d for d in archive_dir.iterdir() if d.is_dir()]
    assert len(subdirs) == 1 and re.match(r"\d{4}\.\d{2}\.\d{2}$", subdirs[0].name)
    archived = {p.name for p in subdirs[0].iterdir()}
    assert "Getsy.xlsx" in archived and "Treolan.xlsx" in archived
    assert any(p.name.startswith("run_report_") for p in subdirs[0].iterdir())
    left = {p.name for p in input_dir.iterdir()}
    assert left == {"Неизвестный прайс.xlsx"}
