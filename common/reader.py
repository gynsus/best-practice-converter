# -*- coding: utf-8 -*-
"""Чтение исходных файлов: xlsx/xlsm (openpyxl), legacy .xls (xlrd).

Ключевой принцип: колонки исходника ищутся ПО ИМЕНАМ заголовков (как в ТЗ V2),
с валидацией. Если поставщик переименовал/сдвинул колонки — воркер падает с
внятной диагностикой, а не пишет смещённые данные.

Дополнительно:
- outline-уровни строк (иерархия групп у Artron, VTT, iDistribute);
- «починка» xlsx с битым styles.xml (кейс Marvel: Excel открывает файл,
  openpyxl — нет; подменяем стили минимальными, данные не затрагиваются).
"""
from __future__ import annotations

import re
import tempfile
import zipfile
from pathlib import Path

import openpyxl


class SourceFormatError(Exception):
    """Структура исходного файла не соответствует ожиданиям воркера."""


_WS = re.compile(r"\s+")

_MIN_STYLES = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    b'<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
    b'<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
    b'<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
    b'<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
    b'<cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>'
    b"</styleSheet>"
)


def repair_styles(path: Path) -> Path:
    """Копия xlsx с заменённым styles.xml (для файлов с битой таблицей стилей)."""
    tmp = Path(tempfile.mktemp(suffix=".xlsx", prefix="repaired_"))
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            zout.writestr(item, _MIN_STYLES if item == "xl/styles.xml" else zin.read(item))
    return tmp


def norm_header(value) -> str:
    """Нормализация имени заголовка: переносы строк/пробелы -> пробел, casefold."""
    if value is None:
        return ""
    return _WS.sub(" ", str(value)).strip().casefold()


class Sheet:
    """Лист исходника: строки значениями + доступ к колонкам по именам заголовков."""

    def __init__(self, path: Path, sheet_name: str, header_row: int,
                 need_outline: bool = False, repair: bool = False):
        self.path = Path(path)
        self.sheet_name = sheet_name
        self.header_row = header_row
        self.outline: dict[int, int] = {}  # номер строки (1-based) -> уровень
        if self.path.suffix.casefold() == ".xls":
            self._load_xls()
        else:
            self._load_xlsx(need_outline=need_outline, repair=repair)
        if len(self.rows) < header_row:
            raise SourceFormatError(f"{self.path.name}: лист «{sheet_name}» короче строки заголовка {header_row}")
        self._header: dict[str, int] = {}
        for idx, v in enumerate(self.rows[header_row - 1]):
            key = norm_header(v)
            if key and key not in self._header:
                self._header[key] = idx

    def _load_xlsx(self, need_outline: bool, repair: bool) -> None:
        src = self.path
        tmp = None
        if repair:
            tmp = src = repair_styles(self.path)
        try:
            # outline-уровни недоступны в read_only-режиме
            wb = openpyxl.load_workbook(src, read_only=not need_outline, data_only=True)
            try:
                if self.sheet_name not in wb.sheetnames:
                    raise SourceFormatError(
                        f"{self.path.name}: нет листа «{self.sheet_name}» (есть: {', '.join(wb.sheetnames)})"
                    )
                ws = wb[self.sheet_name]
                self.rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
                if need_outline:
                    for rn, dim in ws.row_dimensions.items():
                        if dim.outline_level:
                            self.outline[rn] = dim.outline_level
            finally:
                wb.close()
        finally:
            if tmp is not None:
                tmp.unlink(missing_ok=True)

    def _load_xls(self) -> None:
        import xlrd
        wb = xlrd.open_workbook(self.path, formatting_info=True)
        names = wb.sheet_names()
        if self.sheet_name not in names:
            raise SourceFormatError(
                f"{self.path.name}: нет листа «{self.sheet_name}» (есть: {', '.join(names)})"
            )
        sh = wb.sheet_by_name(self.sheet_name)
        self.rows = [
            tuple(None if sh.cell_value(r, c) == "" else sh.cell_value(r, c) for c in range(sh.ncols))
            for r in range(sh.nrows)
        ]
        for r, info in sh.rowinfo_map.items():
            if info.outline_level:
                self.outline[r + 1] = info.outline_level

    def col(self, *names: str) -> int:
        """Индекс (0-based) первой найденной колонки из перечисленных имён.

        Не нашли — SourceFormatError со списком фактических заголовков:
        это штатный механизм обнаружения смены формата поставщиком.
        """
        for name in names:
            idx = self._header.get(norm_header(name))
            if idx is not None:
                return idx
        have = ", ".join(sorted(self._header)) or "<пусто>"
        raise SourceFormatError(
            f"{self.path.name}, лист «{self.sheet_name}»: не найдена колонка "
            f"«{names[0]}» (строка заголовка {self.header_row}; есть: {have})"
        )

    def data_rows(self, first_row: int):
        """Итерация (номер_строки, кортеж значений) начиная с first_row (1-based)."""
        for i in range(first_row - 1, len(self.rows)):
            yield i + 1, self.rows[i]


def cell(row: tuple, idx: int):
    """Безопасное чтение ячейки по индексу (короткие строки в конце листа)."""
    return row[idx] if idx < len(row) else None
